"""Integrity tests for Settings → Data Management restore and xlsx import.

Restore used to commit the "clear existing data" step before reloading the
backup, so any failure mid-restore left the account empty. xlsx import never
passed the owning user to the recipe repository, so every row failed. These
tests use a real file-backed database (not the shared rollback-only fixture)
because the behavior under test is what survives a commit/rollback.
"""

from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.orm import Session, sessionmaker

from app.database.base import Base
from app.dtos.data_management_dtos import (
    DuplicateAction,
    DuplicateResolutionDTO,
    FullBackupDTO,
    RecipeBackupDTO,
)
from app.dtos.unit_conversion_dtos import UnitConversionRuleCreateDTO
from app.dtos.user_category_dtos import UserCategoryCreateDTO, UserCategoryUpdateDTO
from app.dtos.user_ingredient_unit_dtos import UserIngredientUnitCreateDTO
from app.models import (
    Meal,
    NutritionFacts,
    PlannerEntry,
    Recipe,
    RecipeGroup,
    UnitConversionRule,
    User,
    UserCategory,
    UserIngredientUnit,
)
from app.services.data_management import DataManagementService
from app.services.unit_conversion_service import UnitConversionService
from app.services.user_category_service import UserCategoryService
from app.services.user_ingredient_unit_service import UserIngredientUnitService
from app.services.sample_data import SampleDataService


@pytest.fixture()
def file_session(tmp_path) -> Session:
    engine = create_engine(f"sqlite:///{tmp_path / 'integrity.db'}")

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_connection, _record):
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def seeded_user(file_session: Session) -> User:
    user = User(clerk_id="clerk_integrity", email="integrity@example.com", name="I")
    file_session.add(user)
    file_session.commit()
    SampleDataService(file_session, user.id).seed()
    file_session.commit()
    return user


def _count(session: Session, model, user_id: int) -> int:
    return session.execute(
        select(func.count()).select_from(model).where(model.user_id == user_id)
    ).scalar_one()


# ---------------------------------------------------------------------------
# Restore
# ---------------------------------------------------------------------------

class TestRestoreAtomicity:
    def test_failed_restore_keeps_existing_data(self, file_session, seeded_user):
        service = DataManagementService(file_session, seeded_user.id)
        before = {m: _count(file_session, m, seeded_user.id) for m in (Recipe, Meal, PlannerEntry)}
        assert all(before.values())

        backup = service.export_full_backup()
        # Fail partway through the reload, after the clear has run
        with patch(
            "app.services.data_management.restore.RecipeIngredient",
            side_effect=RuntimeError("disk full"),
        ):
            result = service.execute_restore(backup, clear_existing=True)

        assert result.success is False
        assert result.settings is None
        assert "disk full" not in " ".join(result.errors)
        file_session.expire_all()
        after = {m: _count(file_session, m, seeded_user.id) for m in (Recipe, Meal, PlannerEntry)}
        assert after == before

    def test_oversized_meal_tags_are_trimmed_not_fatal(self, file_session, seeded_user):
        service = DataManagementService(file_session, seeded_user.id)
        backup = service.export_full_backup()
        backup.data.meals[0].tags = ["x" * 80] + [f"t{i}" for i in range(30)]

        result = service.execute_restore(backup, clear_existing=True)

        assert result.success is True, result.errors
        tags = [m.tags for m in file_session.query(Meal).filter(Meal.user_id == seeded_user.id)]
        assert all(len(t) <= Meal.MAX_TAGS for t in tags)
        assert all(len(tag) <= Meal.MAX_TAG_LENGTH for t in tags for tag in t)

    def test_round_trip_preserves_recipe_details_groups_and_nutrition(
        self, file_session, seeded_user
    ):
        uid = seeded_user.id
        recipe = file_session.query(Recipe).filter(Recipe.user_id == uid).first()
        recipe.description = "Family favourite"
        recipe.difficulty = "Easy"
        recipe.source_url = "https://example.com/r"
        file_session.add(NutritionFacts(recipe_id=recipe.id, calories=420))
        group = RecipeGroup(name="Weeknight", user_id=uid)
        group.recipes.append(recipe)
        file_session.add(group)
        file_session.commit()
        sample_count = _count(file_session, Recipe, uid)

        service = DataManagementService(file_session, uid)
        backup = FullBackupDTO(**service.export_full_backup().model_dump(mode="json"))
        result = service.execute_restore(backup, clear_existing=True)
        assert result.success is True, result.errors

        restored = (
            file_session.query(Recipe)
            .filter(Recipe.user_id == uid, Recipe.recipe_name == recipe.recipe_name)
            .one()
        )
        assert restored.description == "Family favourite"
        assert restored.difficulty == "Easy"
        assert restored.source_url == "https://example.com/r"
        assert restored.nutrition_facts is not None
        assert restored.nutrition_facts.calories == 420
        # The group survives the clear; membership is re-linked, not duplicated
        groups = file_session.query(RecipeGroup).filter(RecipeGroup.user_id == uid).all()
        assert len(groups) == 1
        assert [r.id for r in groups[0].recipes] == [restored.id]
        # Starter content stays removable after a restore
        assert file_session.query(Recipe).filter(
            Recipe.user_id == uid, Recipe.is_sample.is_(True)
        ).count() == sample_count

    def test_invalid_image_key_is_dropped(self):
        dto = RecipeBackupDTO(
            id=1,
            recipe_name="x",
            recipe_category="y",
            meal_type="Dinner",
            created_at="2026-01-01T00:00:00Z",
            image_key="../../other-folder",
        )
        assert dto.image_key is None


class TestCustomizationBackup:
    """Categories, units, conversion rules and empty groups survive a backup."""

    def _customize(self, session: Session, uid: int) -> None:
        cats = UserCategoryService(session, uid)
        cats.create_category(UserCategoryCreateDTO(label="Grandma's"))
        first_builtin = cats.get_all_categories(include_disabled=True)[0]
        cats.update_category(first_builtin.id, UserCategoryUpdateDTO(is_enabled=False))
        UserIngredientUnitService(session, uid).create_unit(
            UserIngredientUnitCreateDTO(label="Knob")
        )
        UnitConversionService(session, uid).create_rule(
            UnitConversionRuleCreateDTO(
                ingredient_name="Butter", from_unit="tbs", to_unit="stick", factor=0.125
            )
        )
        session.add(RecipeGroup(name="Someday", user_id=uid))
        session.commit()

    def test_restore_into_a_fresh_account_brings_customizations(
        self, file_session, seeded_user, second_user_file
    ):
        self._customize(file_session, seeded_user.id)
        backup = FullBackupDTO(
            **DataManagementService(file_session, seeded_user.id)
            .export_full_backup()
            .model_dump(mode="json")
        )

        uid = second_user_file.id
        result = DataManagementService(file_session, uid).execute_restore(backup)
        assert result.success is True, result.errors

        source_cats = {
            c.value: c
            for c in file_session.query(UserCategory).filter(
                UserCategory.user_id == seeded_user.id
            )
        }
        restored_cats = {
            c.value: c
            for c in file_session.query(UserCategory).filter(UserCategory.user_id == uid)
        }
        assert restored_cats.keys() == source_cats.keys()
        for value, source in source_cats.items():
            assert restored_cats[value].is_enabled == source.is_enabled
            assert restored_cats[value].is_custom == source.is_custom
        assert any(not c.is_enabled for c in restored_cats.values())

        knob = (
            file_session.query(UserIngredientUnit)
            .filter(UserIngredientUnit.user_id == uid, UserIngredientUnit.label == "Knob")
            .one()
        )
        assert knob.is_custom is True
        rule = (
            file_session.query(UnitConversionRule)
            .filter(UnitConversionRule.user_id == uid)
            .one()
        )
        assert (rule.ingredient_name, rule.from_unit, rule.to_unit, rule.factor) == (
            "butter", "tbs", "stick", 0.125
        )
        assert (
            file_session.query(RecipeGroup)
            .filter(RecipeGroup.user_id == uid, RecipeGroup.name == "Someday")
            .count()
            == 1
        )

    def test_same_account_restore_does_not_duplicate(self, file_session, seeded_user):
        uid = seeded_user.id
        self._customize(file_session, uid)
        service = DataManagementService(file_session, uid)
        backup = FullBackupDTO(**service.export_full_backup().model_dump(mode="json"))
        before = [
            _count(file_session, m, uid)
            for m in (UserCategory, UserIngredientUnit, UnitConversionRule, RecipeGroup)
        ]

        result = service.execute_restore(backup, clear_existing=True)

        assert result.success is True, result.errors
        assert before == [
            _count(file_session, m, uid)
            for m in (UserCategory, UserIngredientUnit, UnitConversionRule, RecipeGroup)
        ]

    def test_older_backup_leaves_customizations_alone(self, file_session, seeded_user):
        uid = seeded_user.id
        self._customize(file_session, uid)
        service = DataManagementService(file_session, uid)
        payload = service.export_full_backup().model_dump(mode="json")
        for key in ("recipe_categories", "ingredient_categories", "ingredient_units", "conversion_rules"):
            del payload["data"][key]
        before = _count(file_session, UserCategory, uid)

        result = service.execute_restore(FullBackupDTO(**payload), clear_existing=True)

        assert result.success is True, result.errors
        assert _count(file_session, UserCategory, uid) == before
        assert _count(file_session, UnitConversionRule, uid) == 1


# ---------------------------------------------------------------------------
# Clear all data
# ---------------------------------------------------------------------------

class TestClearAllImages:
    def test_images_are_destroyed_only_after_the_delete_commits(self, file_session, seeded_user):
        uid = seeded_user.id
        recipe = file_session.query(Recipe).filter(Recipe.user_id == uid).first()
        recipe.reference_image_path = (
            "https://res.cloudinary.com/demo/image/upload/v1/"
            f"meal-genie/recipes/{recipe.image_key}/reference.jpg"
        )
        file_session.commit()

        service = DataManagementService(file_session, uid)
        with patch("cloudinary.uploader.destroy") as destroy, patch.object(
            file_session, "commit", side_effect=RuntimeError("boom")
        ):
            with pytest.raises(RuntimeError):
                service.clear_all_data()
        destroy.assert_not_called()


# ---------------------------------------------------------------------------
# xlsx import
# ---------------------------------------------------------------------------

def _workbook(rows: list[tuple[str, str]], ingredients: list[tuple] = ()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Recipes"
    ws.append(["recipe_name", "recipe_category", "meal_type", "servings"])
    for name, category in rows:
        ws.append([name, category, "Dinner", 2])
    ing = wb.create_sheet("Ingredients")
    ing.append(
        ["recipe_name", "recipe_category", "ingredient_name", "ingredient_category", "quantity", "unit"]
    )
    for row in ingredients:
        ing.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestXlsxImport:
    def test_import_creates_recipes_for_the_caller(self, file_session, second_user_file):
        uid = second_user_file.id
        service = DataManagementService(file_session, uid)
        recipes, errors = service.parse_xlsx(
            _workbook(
                [("Soup", "Soups"), ("Stew", "Soups")],
                [("Soup", "Soups", "Carrot", "Produce", 2, "whole")],
            )
        )
        assert not errors

        result = service.execute_import(recipes, [])

        assert result.success is True, result.errors
        assert result.created_count == 2
        assert _count(file_session, Recipe, uid) == 2

    def test_update_resolution_updates_existing_recipe(self, file_session, second_user_file):
        uid = second_user_file.id
        service = DataManagementService(file_session, uid)
        recipes, _ = service.parse_xlsx(_workbook([("Soup", "Soups")]))
        service.execute_import(recipes, [])

        recipes, _ = service.parse_xlsx(
            _workbook([("Soup", "Soups")], [("Soup", "Soups", "Leek", "Produce", 1, "whole")])
        )
        result = service.execute_import(
            recipes,
            [DuplicateResolutionDTO(recipe_name="Soup", recipe_category="Soups", action=DuplicateAction.UPDATE)],
        )

        assert result.updated_count == 1, result.errors
        soup = file_session.query(Recipe).filter(Recipe.user_id == uid).one()
        assert [ri.ingredient.ingredient_name for ri in soup.ingredients] == ["Leek"]

    def test_bad_row_does_not_discard_earlier_rows(self, file_session, second_user_file):
        uid = second_user_file.id
        service = DataManagementService(file_session, uid)
        recipes, _ = service.parse_xlsx(_workbook([("Good", "A"), ("Bad", "B"), ("Also good", "C")]))

        real_create = service._create_recipe

        def flaky(recipe, new_name=None):
            if recipe.recipe_name == "Bad":
                raise RuntimeError("boom")
            return real_create(recipe, new_name=new_name)

        with patch.object(service, "_create_recipe", side_effect=flaky):
            result = service.execute_import(recipes, [])

        assert result.created_count == 2
        assert len(result.errors) == 1
        names = {r.recipe_name for r in file_session.query(Recipe).filter(Recipe.user_id == uid)}
        assert names == {"Good", "Also good"}

    def test_blank_sheet_and_numeric_headers_are_reported_not_crashing(self, file_session, second_user_file):
        wb = Workbook()
        ws = wb.active
        ws.title = "Recipes"
        ws.append([1, 2, 3])
        buf = BytesIO()
        wb.save(buf)
        recipes, errors = DataManagementService(file_session, second_user_file.id).parse_xlsx(
            buf.getvalue()
        )
        assert recipes == []
        assert errors


@pytest.fixture()
def second_user_file(file_session: Session) -> User:
    user = User(clerk_id="clerk_import", email="import@example.com", name="Importer")
    file_session.add(user)
    file_session.commit()
    return user
