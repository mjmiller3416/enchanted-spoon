"""Tenant-isolation tests for the Settings → Data Management surface.

Clear, backup, restore, and xlsx export/import used to run unscoped queries,
so any signed-in user could wipe or download every account's data. These
tests pin each operation to the calling user.
"""

from unittest.mock import patch

import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.main import app
from app.models.ingredient import Ingredient
from app.models.meal import Meal
from app.models.planner_entry import PlannerEntry
from app.models.recipe import Recipe
from app.models.shopping_item import ShoppingItem
from app.models.user import User
from app.services.data_management import DataManagementService
from app.services.sample_data import SampleDataService


CLOUD = "https://res.cloudinary.com/demo/image/upload/v1"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _count(session: Session, model, user_id: int) -> int:
    return session.execute(
        select(func.count()).select_from(model).where(model.user_id == user_id)
    ).scalar_one()


def _snapshot(session: Session, user_id: int) -> dict[str, int]:
    return {
        model.__name__: _count(session, model, user_id)
        for model in (Recipe, Ingredient, Meal, PlannerEntry, ShoppingItem)
    }


def _recipes(session: Session, user_id: int) -> list[Recipe]:
    return list(session.execute(select(Recipe).where(Recipe.user_id == user_id)).scalars().unique())


@pytest.fixture()
def two_seeded_users(db_session: Session, test_user: User, second_user: User) -> tuple[User, User]:
    """Both users get the full starter pack (recipes, meals, planner, shopping)."""
    SampleDataService(db_session, test_user.id).seed()
    SampleDataService(db_session, second_user.id).seed()
    return test_user, second_user


# ---------------------------------------------------------------------------
# Clear all data
# ---------------------------------------------------------------------------

class TestClearAllData:
    def test_only_clears_the_callers_data(self, db_session, two_seeded_users):
        owner, other = two_seeded_users
        other_before = _snapshot(db_session, other.id)
        assert all(other_before.values()), other_before

        with patch("cloudinary.uploader.destroy", return_value={"result": "ok"}):
            DataManagementService(db_session, owner.id).clear_all_data()

        assert not any(_snapshot(db_session, owner.id).values())
        assert _snapshot(db_session, other.id) == other_before

    def test_only_destroys_images_the_recipe_owns(self, db_session, two_seeded_users):
        owner, other = two_seeded_users
        own, starter, foreign = _recipes(db_session, owner.id)[:3]
        victim = _recipes(db_session, other.id)[0]

        own.reference_image_path = (
            f"{CLOUD}/meal-genie/recipes/{own.image_key}/reference_{own.image_key}.jpg"
        )
        starter.reference_image_path = f"{CLOUD}/meal-genie/starter-pack/tacos/reference.jpg"
        # e.g. a restored backup pointing at another account's asset
        foreign.reference_image_path = (
            f"{CLOUD}/meal-genie/recipes/{victim.image_key}/reference_{victim.image_key}.jpg"
        )
        db_session.flush()

        with patch("cloudinary.uploader.destroy", return_value={"result": "ok"}) as destroy:
            counts = DataManagementService(db_session, owner.id).clear_all_data()

        destroy.assert_called_once_with(f"meal-genie/recipes/{own.image_key}/reference_{own.image_key}")
        assert counts["cloudinary_images"] == 1

    def test_requires_a_user(self, db_session):
        with pytest.raises(ValueError):
            DataManagementService(db_session, None).clear_all_data()


# ---------------------------------------------------------------------------
# Full backup / restore
# ---------------------------------------------------------------------------

class TestBackupRestore:
    def test_backup_contains_only_the_callers_data(self, db_session, two_seeded_users):
        owner, other = two_seeded_users
        backup = DataManagementService(db_session, owner.id).export_full_backup()

        owner_recipe_ids = {r.id for r in _recipes(db_session, owner.id)}
        assert {r.id for r in backup.data.recipes} == owner_recipe_ids
        assert {ri.recipe_id for ri in backup.data.recipe_ingredients} <= owner_recipe_ids
        assert len(backup.data.meals) == _count(db_session, Meal, owner.id)
        assert len(backup.data.planner_entries) == _count(db_session, PlannerEntry, owner.id)
        assert len(backup.data.ingredients) == _count(db_session, Ingredient, owner.id)
        assert len(backup.data.shopping_items) == _count(db_session, ShoppingItem, owner.id)

    def test_restore_lands_in_the_callers_account_only(self, db_session, two_seeded_users):
        owner, other = two_seeded_users
        backup = DataManagementService(db_session, owner.id).export_full_backup()
        owner_before = _snapshot(db_session, owner.id)
        other_recipes_before = _count(db_session, Recipe, other.id)

        with patch("cloudinary.uploader.destroy") as destroy:
            result = DataManagementService(db_session, other.id).execute_restore(backup)

        assert result.success, result.errors
        destroy.assert_not_called()
        assert _snapshot(db_session, owner.id) == owner_before
        assert _count(db_session, Recipe, other.id) == len(backup.data.recipes)
        assert _count(db_session, Meal, other.id) == len(backup.data.meals)
        assert other_recipes_before > 0

        # The owner still holds those image_keys, so the restored copies get fresh ones
        owner_keys = {r.image_key for r in _recipes(db_session, owner.id)}
        assert owner_keys.isdisjoint(r.image_key for r in _recipes(db_session, other.id))

    def test_restore_own_backup_keeps_image_keys(self, db_session, two_seeded_users):
        owner, _ = two_seeded_users
        service = DataManagementService(db_session, owner.id)
        backup = service.export_full_backup()
        keys_before = {r.image_key for r in _recipes(db_session, owner.id)}

        with patch("cloudinary.uploader.destroy") as destroy:
            result = service.execute_restore(backup)

        assert result.success, result.errors
        destroy.assert_not_called()
        assert {r.image_key for r in _recipes(db_session, owner.id)} == keys_before

    def test_preview_counts_only_the_callers_recipes(self, db_session, two_seeded_users):
        owner, other = two_seeded_users
        backup = DataManagementService(db_session, owner.id).export_full_backup()
        preview = DataManagementService(db_session, other.id).preview_restore(backup)

        expected = f"Existing data ({_count(db_session, Recipe, other.id)} recipes)"
        assert any(w.startswith(expected) for w in preview.warnings), preview.warnings


# ---------------------------------------------------------------------------
# xlsx export / import
# ---------------------------------------------------------------------------

class TestXlsxExportImport:
    def test_export_contains_only_the_callers_recipes(self, db_session, two_seeded_users):
        from io import BytesIO

        from openpyxl import load_workbook

        owner, other = two_seeded_users
        other_recipe = _recipes(db_session, other.id)[0]
        other_recipe.recipe_name = "Other User Secret Recipe"
        db_session.flush()

        data = DataManagementService(db_session, owner.id).export_recipes_to_xlsx()
        sheet = load_workbook(BytesIO(data))["Recipes"]
        names = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]

        assert len(names) == _count(db_session, Recipe, owner.id)
        assert "Other User Secret Recipe" not in names

    def test_import_duplicate_check_ignores_other_users(self, db_session, two_seeded_users):
        owner, other = two_seeded_users
        other_recipe = _recipes(db_session, other.id)[0]
        other_recipe.recipe_name = "Only In Other Account"
        db_session.flush()

        service = DataManagementService(db_session, owner.id)
        assert service._find_existing_recipe("Only In Other Account", other_recipe.recipe_category) is None


# ---------------------------------------------------------------------------
# Shopping ingredient breakdown
# ---------------------------------------------------------------------------

def test_ingredient_breakdown_ignores_other_users_recipes(db_session, two_seeded_users):
    from app.services.shopping import ShoppingService

    owner, other = two_seeded_users
    own_recipe = _recipes(db_session, owner.id)[0]
    other_recipe = _recipes(db_session, other.id)[0]
    other_recipe.recipe_name = "Other User Secret Recipe"
    db_session.flush()

    breakdown = ShoppingService(db_session, owner.id).get_ingredient_breakdown(
        [own_recipe.id, other_recipe.id]
    )

    recipe_names = {
        row.recipe_name for item in breakdown.items for row in item.recipe_breakdown
    }
    assert recipe_names == {own_recipe.recipe_name}


# ---------------------------------------------------------------------------
# Upload routes
# ---------------------------------------------------------------------------

def test_arbitrary_cloudinary_delete_route_is_gone():
    paths = app.openapi()["paths"]
    assert any(path.startswith("/api/upload") for path in paths)
    upload_deletes = [
        path for path, ops in paths.items() if path.startswith("/api/upload") and "delete" in ops
    ]
    assert upload_deletes == []
